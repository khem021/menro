"""
Generate MENRO Post-Deployment Setup Guide as a .docx Word document.
Run: python make_postdeploy_guide.py
Output: MENRO_Post_Deployment_Setup_Guide.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ── Helpers ───────────────────────────────────────────────────────────────────
def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_table_borders(table, color='4CAF50', size='4'):
    tbl  = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    size)
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        tblBorders.append(b)
    tblPr.append(tblBorders)

def add_code(doc, text, color='9CD377'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  '1E1E1E')
    pPr.append(shd)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '360')
    pPr.append(ind)
    run = p.add_run(text)
    run.font.name  = 'Courier New'
    run.font.size  = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    return p

def add_code_comment(doc, text):
    return add_code(doc, text, color='6A9955')

def add_tip(doc, icon, label, text, bg='E8F5E9', label_color='1B5E20'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after  = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  bg)
    pPr.append(shd)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'),  '280')
    ind.set(qn('w:right'), '280')
    pPr.append(ind)
    r1 = p.add_run(f'{icon}  {label}  ')
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor.from_string(label_color)
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    return p

def add_body(doc, text, bold=False, size=10.5, color=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def add_bullet(doc, text, level=0, bold_part=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.left_indent  = Inches(0.3 * (level + 1))
    if bold_part and text.startswith(bold_part):
        r1 = p.add_run(bold_part)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r2 = p.add_run(text[len(bold_part):])
        r2.font.size = Pt(10)
    else:
        run = p.add_run(text)
        run.font.size = Pt(10)
    return p

def add_num(doc, text, level=0):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.left_indent  = Inches(0.3 * (level + 1))
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    return p

def add_section_title(doc, text, size=15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after  = Pt(6)
    # underline rule
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '4CAF50')
    pBdr.append(bottom)
    p._p.get_or_add_pPr().append(pBdr)
    run = p.add_run(text)
    run.font.size  = Pt(size)
    run.font.bold  = True
    run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
    return p

def add_step_badge(doc, number, title, subtitle=''):
    # Two-cell row: big number badge + title
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.columns[0].width = Cm(1.8)
    tbl.columns[1].width = Cm(13)
    tbl.rows[0].height = Cm(1.4)

    c0 = tbl.rows[0].cells[0]
    c1 = tbl.rows[0].cells[1]

    shade_cell(c0, '1B5E20')
    shade_cell(c1, '2E7D32')

    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after  = Pt(0)
    r0 = p0.add_run(str(number))
    r0.font.size  = Pt(26)
    r0.font.bold  = True
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after  = Pt(0)
    p1.paragraph_format.left_indent  = Cm(0.4)
    r1 = p1.add_run(title + '\n')
    r1.font.size  = Pt(14)
    r1.font.bold  = True
    r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if subtitle:
        r2 = p1.add_run(subtitle)
        r2.font.size  = Pt(9)
        r2.font.color.rgb = RGBColor(0xA5, 0xD6, 0xA7)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


# ═══════════════════════════════════════════════════════════════════════════════
#  COVER
# ═══════════════════════════════════════════════════════════════════════════════

# Top accent bar
p_bar = doc.add_paragraph()
p_bar.paragraph_format.space_before = Pt(0)
p_bar.paragraph_format.space_after  = Pt(0)
pPr = p_bar._p.get_or_add_pPr()
shd = OxmlElement('w:shd')
shd.set(qn('w:val'),   'clear')
shd.set(qn('w:color'), 'auto')
shd.set(qn('w:fill'),  '4CAF50')
pPr.append(shd)
run_bar = p_bar.add_run(' ')
run_bar.font.size = Pt(6)

# Main title block
p_main = doc.add_paragraph()
p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_main.paragraph_format.space_before = Pt(50)
p_main.paragraph_format.space_after  = Pt(0)
pPr2 = p_main._p.get_or_add_pPr()
shd2 = OxmlElement('w:shd')
shd2.set(qn('w:val'),   'clear')
shd2.set(qn('w:color'), 'auto')
shd2.set(qn('w:fill'),  '1B5E20')
pPr2.append(shd2)
r_title = p_main.add_run('  POST-DEPLOYMENT SETUP  ')
r_title.font.size  = Pt(32)
r_title.font.bold  = True
r_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
r_title.font.name  = 'Calibri'

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_before = Pt(0)
p_sub.paragraph_format.space_after  = Pt(0)
pPr3 = p_sub._p.get_or_add_pPr()
shd3 = OxmlElement('w:shd')
shd3.set(qn('w:val'),   'clear')
shd3.set(qn('w:color'), 'auto')
shd3.set(qn('w:fill'),  '2E7D32')
pPr3.append(shd3)
r_sub = p_sub.add_run('  MENRO Waste Management Information System  ')
r_sub.font.size  = Pt(13)
r_sub.font.color.rgb = RGBColor(0xC8, 0xE6, 0xC9)
r_sub.font.name  = 'Calibri'

doc.add_paragraph()

# "What this guide covers" intro box
tbl_intro = doc.add_table(rows=1, cols=1)
tbl_intro.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_intro.columns[0].width = Cm(13)
set_table_borders(tbl_intro, 'A5D6A7', '6')
c_intro = tbl_intro.rows[0].cells[0]
shade_cell(c_intro, 'F1F8E9')
p_intro = c_intro.paragraphs[0]
p_intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_intro.paragraph_format.space_before = Pt(14)
p_intro.paragraph_format.space_after  = Pt(14)
r_i1 = p_intro.add_run('What this guide covers\n')
r_i1.font.bold = True; r_i1.font.size = Pt(12); r_i1.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
r_i2 = p_intro.add_run(
    'Two essential tasks you must complete immediately after\n'
    'deploying the MENRO system on your Hostinger VPS.'
)
r_i2.font.size = Pt(10.5)

doc.add_paragraph()

# Two step previews
tbl_prev = doc.add_table(rows=1, cols=2)
tbl_prev.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_prev.columns[0].width = Cm(7)
tbl_prev.columns[1].width = Cm(7)
set_table_borders(tbl_prev, 'C8E6C9', '4')

steps_preview = [
    ('1', 'Build the\nFrontend Assets', 'Run npm run build to compile\nCSS, JavaScript and images.', 'E8F5E9'),
    ('2', 'Change All\nDefault Passwords', 'Secure the system by updating\nall 6 default user accounts.', 'F1F8E9'),
]
for cell, (num, title, desc, bg) in zip(tbl_prev.rows[0].cells, steps_preview):
    shade_cell(cell, bg)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(14)
    rn = p.add_run(num + '\n')
    rn.font.size = Pt(30); rn.font.bold = True; rn.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
    rt = p.add_run(title + '\n')
    rt.font.size = Pt(11); rt.font.bold = True; rt.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
    rd = p.add_run(desc)
    rd.font.size = Pt(9); rd.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

add_tip(doc, '⚡', 'Time required:',
        'About 5–10 minutes total.',
        bg='FFF8E1', label_color='E65100')

p_date = doc.add_paragraph()
p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_date.paragraph_format.space_before = Pt(30)
r_d = p_date.add_run('MENRO Waste Management System  ·  Post-Deployment Guide  ·  v1.0  ·  May 2026')
r_d.font.size = Pt(8); r_d.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  BEFORE YOU BEGIN
# ═══════════════════════════════════════════════════════════════════════════════
add_section_title(doc, 'Before You Begin')

add_body(doc, 'Make sure you have already completed the full server deployment from the Hostinger Deployment Guide. At this point you should have:')
add_bullet(doc, 'The MENRO project files uploaded to /var/www/menro on your VPS')
add_bullet(doc, 'The deploy.sh script already run successfully')
add_bullet(doc, 'MySQL database created and .env file configured')
add_bullet(doc, 'Nginx running and pointing to your domain')
doc.add_paragraph()
add_tip(doc, '💻', 'How to connect:',
        'All commands below are run on the SERVER via SSH.\n'
        'Open your terminal and connect first:  ssh root@YOUR_VPS_IP',
        bg='E3F2FD', label_color='0D47A1')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  STEP 1 — BUILD FRONTEND
# ═══════════════════════════════════════════════════════════════════════════════
add_step_badge(doc, 1, 'Build the Frontend Assets',
               'Compile CSS, JavaScript, and image files for production')

add_body(doc, 'Why is this needed?', bold=True)
add_body(doc,
    'The MENRO system uses Vite to bundle all CSS (Tailwind) and JavaScript (Alpine.js) '
    'into optimized files. These bundled files live in the public/build/ folder and are '
    'NOT stored in Git (they are in .gitignore). Without running this step, the site will '
    'load but will have no styling — it will appear as a plain, unstyled HTML page.')

doc.add_paragraph()
add_tip(doc, '⚠', 'What you will see if you skip this:',
        'The site loads but looks completely broken — no colors, no layout, no icons. '
        'This is normal and fixed immediately by running the command below.',
        bg='FFF3CD', label_color='856404')

doc.add_paragraph()
add_body(doc, 'Step-by-step:', bold=True)
add_num(doc, 'SSH into your VPS (if not already connected):')
add_code(doc, 'ssh root@YOUR_VPS_IP')
doc.add_paragraph()

add_num(doc, 'Go to the app directory:')
add_code(doc, 'cd /var/www/menro')
doc.add_paragraph()

add_num(doc, 'Install Node packages:')
add_code(doc, 'npm ci --ignore-scripts')
add_code_comment(doc, '# This reads package-lock.json and installs exact versions')
add_code_comment(doc, '# Takes about 1–2 minutes on first run')
doc.add_paragraph()

add_num(doc, 'Build the production assets:')
add_code(doc, 'npm run build')
add_code_comment(doc, '# Compiles Tailwind CSS + Alpine.js into public/build/')
add_code_comment(doc, '# Takes about 30–60 seconds')
doc.add_paragraph()

add_num(doc, 'Clean up (removes ~300 MB of dev packages no longer needed):')
add_code(doc, 'rm -rf node_modules')
doc.add_paragraph()

add_num(doc, 'Restart PHP-FPM so Laravel picks up all changes:')
add_code(doc, 'systemctl restart php8.3-fpm')
add_code(doc, 'systemctl reload nginx')
doc.add_paragraph()

add_body(doc, 'What a successful build looks like:', bold=True)
add_code(doc, 'vite v4.x.x building for production...')
add_code(doc, '')
add_code(doc, '✓  530 modules transformed.')
add_code(doc, 'public/build/manifest.json             1.23 kB')
add_code(doc, 'public/build/assets/app-Bx9K2m4d.css   98.34 kB │ gzip: 16.20 kB')
add_code(doc, 'public/build/assets/app-DjFk9P2l.js   142.78 kB │ gzip: 45.66 kB')
add_code(doc, '✓  Built in 42.31s')
doc.add_paragraph()

add_tip(doc, '✅', 'Done!',
        'Open your browser and visit https://your-domain.com — the site should now display '
        'correctly with full styling, colors, and layout.',
        bg='E8F5E9', label_color='1B5E20')

# ── Troubleshooting ───────────────────────────────────────────────────────────
add_body(doc, 'Troubleshooting npm run build', bold=True, size=11, color='1B5E20', space_after=4)

issues_npm = [
    ('node: command not found',
     'Node.js is not installed. Run:\n'
     'curl -fsSL https://deb.nodesource.com/setup_20.x | bash -\n'
     'apt-get install -y nodejs'),
    ('npm: command not found',
     'npm is bundled with Node.js. Reinstall Node.js using the command above.'),
    ('EACCES permission denied',
     'Run the command as root or with sudo. You should already be root on the VPS.'),
    ('JavaScript heap out of memory',
     'Increase Node memory limit:\n'
     'NODE_OPTIONS=--max-old-space-size=512 npm run build'),
    ('vite: not found  /  missing build script',
     'node_modules was deleted before building. Run npm ci first, then npm run build.'),
]
tbl_npm = doc.add_table(rows=1 + len(issues_npm), cols=2)
tbl_npm.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl_npm.columns[0].width = Cm(6)
tbl_npm.columns[1].width = Cm(9.5)
set_table_borders(tbl_npm, '4CAF50', '4')
for c, t in zip(tbl_npm.rows[0].cells, ['Error / Problem', 'Fix']):
    shade_cell(c, '2E7D32')
    pp = c.paragraphs[0]
    pp.paragraph_format.space_before = Pt(4); pp.paragraph_format.space_after = Pt(4)
    rr = pp.add_run(t); rr.font.bold = True; rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for i, (err, fix) in enumerate(issues_npm):
    bg = 'F9FBE7' if i % 2 == 0 else 'FFFFFF'
    c0 = tbl_npm.rows[i+1].cells[0]; c1 = tbl_npm.rows[i+1].cells[1]
    shade_cell(c0, bg); shade_cell(c1, bg)
    for c, t in [(c0, err), (c1, fix)]:
        pp = c.paragraphs[0]
        pp.paragraph_format.space_before = Pt(4); pp.paragraph_format.space_after = Pt(4)
        rr = pp.add_run(t); rr.font.size = Pt(9); rr.font.name = 'Courier New' if c == c0 else 'Calibri'

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  STEP 2 — CHANGE PASSWORDS
# ═══════════════════════════════════════════════════════════════════════════════
add_step_badge(doc, 2, 'Change All Default Passwords',
               'Secure the system — do this before sharing the URL with anyone')

add_body(doc, 'Why is this critical?', bold=True)
add_body(doc,
    'The system ships with 6 default user accounts, all using the password admin123. '
    'These credentials are public knowledge (they are documented in this guide and in the '
    'source code). Anyone who finds your site URL can log in immediately if you do not '
    'change the passwords. This must be done before sharing the system with any staff or users.')
doc.add_paragraph()
add_tip(doc, '🔴', 'Security Warning:',
        'Do NOT skip this step. Default passwords are a critical security risk. '
        'Change every password listed below before going live.',
        bg='FFEBEE', label_color='B71C1C')

doc.add_paragraph()
add_body(doc, 'Default accounts that must be updated:', bold=True)

accounts = [
    ('admin',     'admin123', 'System Administrator', 'Full system access — highest priority to change',   'B71C1C'),
    ('menro',     'admin123', 'MENRO Officer',         'Full operational access',                           'B71C1C'),
    ('encoder',   'admin123', 'Data Encoder',          'Can create waste entries and generator records',    'E65100'),
    ('inspector', 'admin123', 'Field Inspector',       'Can create inspections and violations',             'E65100'),
    ('barangay',  'admin123', 'Barangay User',         'Can view barangay data',                            '1B5E20'),
    ('viewer',    'admin123', 'Report Viewer',         'Read-only access to reports',                       '1B5E20'),
]

tbl_acc = doc.add_table(rows=1 + len(accounts), cols=5)
tbl_acc.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl_acc.columns[0].width = Cm(2.8)
tbl_acc.columns[1].width = Cm(2.5)
tbl_acc.columns[2].width = Cm(3.5)
tbl_acc.columns[3].width = Cm(5.5)
tbl_acc.columns[4].width = Cm(1.4)
set_table_borders(tbl_acc, '4CAF50', '4')

for c, t in zip(tbl_acc.rows[0].cells, ['Username', 'Default Password', 'Role', 'Access Level', '✓ Done']):
    shade_cell(c, '2E7D32')
    pp = c.paragraphs[0]
    pp.paragraph_format.space_before = Pt(4); pp.paragraph_format.space_after = Pt(4)
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = pp.add_run(t); rr.font.bold = True; rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for i, (user, pw, role, access, priority_color) in enumerate(accounts):
    bg = 'FFEBEE' if i < 2 else ('FFF3E0' if i < 4 else 'F1F8E9')
    cells = tbl_acc.rows[i+1].cells
    for c in cells: shade_cell(c, bg)
    for c, t, mono in [
        (cells[0], user,   True),
        (cells[1], pw,     True),
        (cells[2], role,   False),
        (cells[3], access, False),
        (cells[4], '☐',   False),
    ]:
        pp = c.paragraphs[0]
        pp.paragraph_format.space_before = Pt(5)
        pp.paragraph_format.space_after  = Pt(5)
        if c == cells[4]: pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(t)
        rr.font.size = Pt(9.5)
        if mono: rr.font.name = 'Courier New'; rr.font.bold = True
        if c == cells[0]: rr.font.color.rgb = RGBColor.from_string(priority_color)
        if c == cells[4]: rr.font.size = Pt(13)

doc.add_paragraph()

# ── How to change password ────────────────────────────────────────────────────
add_body(doc, 'How to change a password — step by step:', bold=True, size=11, color='1B5E20')
add_num(doc, 'Open your browser and go to:  https://your-domain.com')
add_num(doc, 'Log in using one of the accounts above (e.g. username: admin, password: admin123).')
add_num(doc, 'Once logged in, click the user avatar / name in the top-right corner of the screen.')
add_num(doc, 'Select "Settings" or "My Profile" from the dropdown menu.')
add_num(doc, 'Find the Change Password section.')
add_num(doc, 'Enter the current password (admin123), then enter a new strong password twice.')
add_num(doc, 'Click Save / Update Password.')
add_num(doc, 'Log out, then log back in with the new password to confirm it works.')
add_num(doc, 'Repeat for every account in the table above.')
doc.add_paragraph()

# ── Admin method ──────────────────────────────────────────────────────────────
add_body(doc, 'Faster method — change all passwords from the Admin account:', bold=True, size=11, color='1B5E20')
add_body(doc,
    'The System Administrator account (admin) can edit all other users directly '
    'from the User Management panel:')
add_num(doc, 'Log in as admin.')
add_num(doc, 'Go to the sidebar → Users.')
add_num(doc, 'Click Edit on each user.')
add_num(doc, 'Set a new password in the Password field.')
add_num(doc, 'Click Save.')
add_num(doc, 'Repeat for all 6 users including the admin account itself.')
doc.add_paragraph()

# ── Strong password tips ──────────────────────────────────────────────────────
add_body(doc, 'Password requirements and best practices:', bold=True, size=11, color='1B5E20')

tbl_pw = doc.add_table(rows=1, cols=2)
tbl_pw.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl_pw.columns[0].width = Cm(7)
tbl_pw.columns[1].width = Cm(8.5)
set_table_borders(tbl_pw, 'A5D6A7', '4')
shade_cell(tbl_pw.rows[0].cells[0], 'E8F5E9')
shade_cell(tbl_pw.rows[0].cells[1], 'FFEBEE')

pp_l = tbl_pw.rows[0].cells[0].paragraphs[0]
pp_l.paragraph_format.space_before = Pt(8)
pp_l.paragraph_format.space_after  = Pt(8)
pp_l.paragraph_format.left_indent  = Cm(0.3)
r_l = pp_l.add_run('✅  Good password examples\n\n')
r_l.font.bold = True; r_l.font.size = Pt(10); r_l.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
for good in ['Menro@2026!', 'Madrid#Waste99', 'Env!r0nment2026']:
    rg = pp_l.add_run(f'   {good}\n')
    rg.font.name = 'Courier New'; rg.font.size = Pt(9.5)
    rg.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)

pp_r = tbl_pw.rows[0].cells[1].paragraphs[0]
pp_r.paragraph_format.space_before = Pt(8)
pp_r.paragraph_format.space_after  = Pt(8)
pp_r.paragraph_format.left_indent  = Cm(0.3)
r_r = pp_r.add_run('❌  Avoid these patterns\n\n')
r_r.font.bold = True; r_r.font.size = Pt(10); r_r.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)
for bad, reason in [('admin123', 'Default — everyone knows it'), ('password', 'Too common'), ('menro2024', 'Too guessable')]:
    rb = pp_r.add_run(f'   {bad:<14} ')
    rb.font.name = 'Courier New'; rb.font.size = Pt(9.5)
    rb.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)
    rr2 = pp_r.add_run(f'({reason})\n')
    rr2.font.size = Pt(9); rr2.font.italic = True
    rr2.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

doc.add_paragraph()
add_tip(doc, '🔒', 'Minimum recommendation:',
        'At least 10 characters, mix of uppercase, lowercase, numbers, and a symbol (!, @, #, etc.).',
        bg='E8F5E9', label_color='1B5E20')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  VERIFICATION CHECKLIST
# ═══════════════════════════════════════════════════════════════════════════════
add_section_title(doc, 'Final Verification Checklist')
add_body(doc, 'Use this checklist to confirm the system is fully set up and secured:')
doc.add_paragraph()

checks = [
    ('Frontend Build', [
        ('public/build/ folder exists on the server', False),
        ('Site loads with full styling and colors (no plain HTML)', False),
        ('Login page displays correctly with logo images', False),
        ('Dashboard loads after logging in', False),
    ]),
    ('Password Security', [
        ('admin account password changed from admin123', False),
        ('menro account password changed from admin123', False),
        ('encoder account password changed from admin123', False),
        ('inspector account password changed from admin123', False),
        ('barangay account password changed from admin123', False),
        ('viewer account password changed from admin123', False),
    ]),
    ('System Functions', [
        ('Can create a new waste generator', False),
        ('Can add a waste entry', False),
        ('Can generate and download a report (.xlsx)', False),
        ('Analytics page loads with charts', False),
    ]),
    ('Security', [
        ('APP_DEBUG=false in .env', False),
        ('APP_ENV=production in .env', False),
        ('HTTPS is enabled (padlock visible in browser)', False),
        ('Accessing http:// redirects to https://', False),
    ]),
]

for group_title, items in checks:
    # group header
    ph = doc.add_paragraph()
    ph.paragraph_format.space_before = Pt(8)
    ph.paragraph_format.space_after  = Pt(3)
    pPr_h = ph._p.get_or_add_pPr()
    shd_h = OxmlElement('w:shd')
    shd_h.set(qn('w:val'),   'clear')
    shd_h.set(qn('w:color'), 'auto')
    shd_h.set(qn('w:fill'),  'E8F5E9')
    pPr_h.append(shd_h)
    ind_h = OxmlElement('w:ind')
    ind_h.set(qn('w:left'), '200')
    pPr_h.append(ind_h)
    rh = ph.add_run(f'  {group_title}')
    rh.font.bold = True; rh.font.size = Pt(11)
    rh.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    for label, checked in items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after  = Pt(2)
        p_item.paragraph_format.left_indent  = Cm(0.8)
        r_box = p_item.add_run('☐   ')
        r_box.font.size = Pt(12)
        r_box.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
        r_label = p_item.add_run(label)
        r_label.font.size = Pt(10.5)

doc.add_paragraph()
add_tip(doc, '🎉', 'You are done!',
        'Once all items above are checked, the MENRO system is fully deployed, '
        'secured, and ready for daily use by your staff.',
        bg='E8F5E9', label_color='1B5E20')

# ═══════════════════════════════════════════════════════════════════════════════
#  QUICK REFERENCE — server commands
# ═══════════════════════════════════════════════════════════════════════════════
add_section_title(doc, 'Quick Reference — Useful Server Commands')
add_body(doc, 'These are the most common commands you may need after deployment:')
doc.add_paragraph()

cmds = [
    ('Rebuild frontend after code update',
     'cd /var/www/menro\nnpm ci --ignore-scripts && npm run build && rm -rf node_modules'),
    ('Clear all Laravel caches',
     'cd /var/www/menro\nphp artisan config:clear\nphp artisan view:clear\nphp artisan cache:clear'),
    ('Rebuild caches for production',
     'cd /var/www/menro\nphp artisan config:cache\nphp artisan view:cache'),
    ('Restart web services',
     'systemctl restart php8.3-fpm\nsystemctl reload nginx'),
    ('View Laravel error log',
     'tail -n 50 /var/www/menro/storage/logs/laravel.log'),
    ('Check Nginx status / errors',
     'systemctl status nginx\njournalctl -u nginx --since "5 min ago"'),
    ('Check PHP-FPM status',
     'systemctl status php8.3-fpm'),
    ('Fix permissions (if 403 errors)',
     'chown -R www-data:www-data /var/www/menro\nchmod -R 775 /var/www/menro/storage /var/www/menro/bootstrap/cache'),
]
for cmd_title, cmd_text in cmds:
    add_body(doc, cmd_title, bold=True, size=10, color='1B5E20', space_after=1)
    for line in cmd_text.split('\n'):
        add_code(doc, line)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ── Footer ────────────────────────────────────────────────────────────────────
doc.add_paragraph()
p_foot = doc.add_paragraph()
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_foot.paragraph_format.space_before = Pt(20)
r_foot = p_foot.add_run(
    'MENRO Waste Management Information System  ·  Post-Deployment Setup Guide  ·  v1.0  ·  2026'
)
r_foot.font.size = Pt(8); r_foot.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r'C:\Users\jhonr\Herd\menro\MENRO_Post_Deployment_Setup_Guide.docx'
doc.save(out)
print(f'Saved: {out}')
