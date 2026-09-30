"""
Generate the MENRO ISO/IEC 25010:2023 evaluation questionnaire as a .docx.
Run: python make_iso25010_questionnaire.py
Output: MENRO_ISO25010_Evaluation_Questionnaire.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = '2E7D32'
LIGHT = 'E8F5E9'

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.59), Cm(33.02)   # Long bond (8.5" x 13")
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.0)

st = doc.styles['Normal']
st.font.name = 'Arial'
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def borders(table, color='9E9E9E'):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{name}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    h.set(qn('w:val'), 'true')
    trPr.append(h)


def keep_row(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement('w:cantSplit')
    c.set(qn('w:val'), 'true')
    trPr.append(c)


def cell_text(cell, text, bold=False, size=10, align=None, color=None, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def para(text='', bold=False, size=10, align=None, after=4, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(text):
    p = para(text, bold=True, size=11, after=3, color=GREEN)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    return p


def set_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


# ── Title block ──────────────────────────────────────────────────────────────
para('Republic of the Philippines', size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para('Municipality of Madrid, Surigao del Sur', size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para('Municipal Environment and Natural Resources Office (MENRO)', size=9,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para('SYSTEM EVALUATION QUESTIONNAIRE', bold=True, size=14,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=0, color=GREEN)
para('MENRO Waste Management System', bold=True, size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para('Based on ISO/IEC 25010:2023 — Product Quality Model', italic=True, size=10,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

# ── Part I: Respondent profile ───────────────────────────────────────────────
heading('PART I. RESPONDENT PROFILE')
prof = doc.add_table(rows=4, cols=2)
borders(prof)
rows = [
    ('Name (optional):', ''),
    ('Office / Barangay:', ''),
    ('Role in the system:',
     '☐ System Administrator   ☐ MENRO Officer   ☐ Data Encoder\n'
     '☐ Field Inspector   ☐ Barangay User   ☐ Report Viewer   ☐ IT Expert'),
    ('Date of evaluation:', ''),
]
for i, (k, v) in enumerate(rows):
    cell_text(prof.rows[i].cells[0], k, bold=True)
    shade(prof.rows[i].cells[0], LIGHT)
    prof.rows[i].cells[1].text = ''
    lines = v.split('\n')
    p = prof.rows[i].cells[1].paragraphs[0]
    for j, line in enumerate(lines):
        if j:
            p = prof.rows[i].cells[1].add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.add_run(line).font.size = Pt(9.5)
set_widths(prof, [4.5, 13.0])

# ── Instructions & scale ─────────────────────────────────────────────────────
heading('PART II. INSTRUCTIONS')
para('After using the MENRO Waste Management System, rate each statement by placing a '
     'check (✓) in the column that best matches your experience. Answer every item honestly; '
     'there are no right or wrong answers. Items marked with an asterisk (*) are intended for '
     'IT experts / technical evaluators; other respondents may leave them blank.', after=6)

scale = doc.add_table(rows=6, cols=3)
borders(scale)
scale.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ('Scale', 'Verbal Description', 'Weighted Mean Range')
for i, h in enumerate(hdr):
    cell_text(scale.rows[0].cells[i], h, bold=True, color='FFFFFF', align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(scale.rows[0].cells[i], GREEN)
for r, (n, d, rng) in enumerate([
    ('5', 'Strongly Agree (Excellent)', '4.21 – 5.00'),
    ('4', 'Agree (Very Good)', '3.41 – 4.20'),
    ('3', 'Neutral (Good)', '2.61 – 3.40'),
    ('2', 'Disagree (Fair)', '1.81 – 2.60'),
    ('1', 'Strongly Disagree (Poor)', '1.00 – 1.80'),
], start=1):
    cell_text(scale.rows[r].cells[0], n, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(scale.rows[r].cells[1], d)
    cell_text(scale.rows[r].cells[2], rng, align=WD_ALIGN_PARAGRAPH.CENTER)
set_widths(scale, [2.0, 6.0, 4.5])

# ── Part III: Questionnaire items ────────────────────────────────────────────
# (characteristic, definition, [(sub-characteristic, statement, technical?)])
SECTIONS = [
    ('A. FUNCTIONAL SUITABILITY',
     'Degree to which the system provides functions that meet the needs of MENRO operations.',
     [
         ('Functional completeness',
          'The system covers all core MENRO tasks: waste generators, waste entries, collections, '
          'inspections, violations, incidents, and reports.', False),
         ('Functional completeness',
          'The barangay, sector, and cluster records support how waste is managed per area '
          'in the municipality.', False),
         ('Functional correctness',
          'Totals of waste quantities (by category, generator type, barangay, and cluster) '
          'on the dashboard and analytics are accurate.', False),
         ('Functional correctness',
          'Generator compliance status updates correctly based on inspection and violation records.', False),
         ('Functional correctness',
          'Exported and archived reports (Excel/PDF) contain the same data shown in the system.', False),
         ('Functional appropriateness',
          'Automated features (collection reminders, overdue follow-up flags, and scheduled '
          'daily/weekly/monthly/yearly reports) reduce manual work.', False),
     ]),
    ('B. PERFORMANCE EFFICIENCY',
     'Performance relative to the amount of resources used under stated conditions.',
     [
         ('Time behaviour', 'Pages such as the dashboard, lists, and forms load quickly.', False),
         ('Time behaviour',
          'Searching, filtering, and saving records respond without noticeable delay.', False),
         ('Time behaviour', 'Generating, exporting, or printing reports finishes in an acceptable time.', False),
         ('Resource utilization',
          'The system runs smoothly on the office computers and internet connection available at MENRO.', False),
         ('Capacity',
          'The system stays responsive as records grow and several users work at the same time.', False),
     ]),
    ('C. COMPATIBILITY',
     'Degree to which the system can exchange information and operate alongside other products.',
     [
         ('Co-existence',
          'The system works properly while other applications or browser tabs are open.', False),
         ('Interoperability',
          'Exported Excel and PDF files open correctly in common programs (e.g., MS Excel, PDF readers).', False),
         ('Interoperability',
          'The system works consistently on common web browsers (Chrome, Edge, Firefox).', False),
     ]),
    ('D. INTERACTION CAPABILITY (formerly Usability)',
     'Degree to which users can interact with the system effectively, efficiently, and with satisfaction.',
     [
         ('Appropriateness recognizability',
          'I can easily tell from the menus which module to use for each MENRO task.', False),
         ('Learnability', 'I learned to use the system quickly with little or no training.', False),
         ('Operability',
          'Forms for encoding entries, inspections, violations, and incidents are easy to fill out.', False),
         ('User error protection',
          'The system validates inputs and asks for confirmation before deleting records.', False),
         ('User engagement', 'The interface is clean, consistent, and pleasant to use.', False),
         ('Inclusivity',
          'The system is usable on different screen sizes (desktop, laptop, tablet, phone).', False),
         ('User assistance / Self-descriptiveness',
          'Labels, messages, and notifications clearly explain what is happening and what to do next.', False),
     ]),
    ('E. RELIABILITY',
     'Degree to which the system performs its functions without failure over a period of time.',
     [
         ('Faultlessness', 'The system operates without errors or crashes during normal use.', False),
         ('Availability', 'The system is accessible whenever I need to use it.', False),
         ('Fault tolerance',
          'Invalid input or a lost connection does not corrupt or lose saved records.', False),
         ('Recoverability',
          'Data can be recovered after a failure, and archived reports remain available for download.', False),
     ]),
    ('F. SECURITY',
     'Degree to which the system protects information and data from unauthorized access or modification.',
     [
         ('Confidentiality',
          'Users can only access modules allowed for their role (e.g., only admins manage users and the archive).', False),
         ('Integrity', 'Records cannot be changed by users who are not authorized to do so.', False),
         ('Accountability / Non-repudiation',
          'The audit log records who performed each action and when.', False),
         ('Authenticity',
          'The login protects accounts using a username and password, and users can change their own password.', False),
         ('Resistance',
          'The system resists common attacks (e.g., repeated login attempts, SQL injection, CSRF).*', True),
     ]),
    ('G. MAINTAINABILITY',
     'Degree to which the system can be modified effectively and efficiently by developers.',
     [
         ('Modularity',
          'The system is divided into separate modules, so changing one does not break others.*', True),
         ('Reusability',
          'Components (e.g., report builder, Livewire forms) can be reused for new features.*', True),
         ('Analysability', 'Errors can be traced easily through logs and the audit trail.*', True),
         ('Modifiability',
          'Barangay, cluster, and user records can be updated through the system without changing the code.', False),
         ('Testability', 'The system has automated tests that verify its main workflows.*', True),
     ]),
    ('H. FLEXIBILITY (formerly Portability)',
     'Degree to which the system can be adapted to changes in requirements, contexts, or environments.',
     [
         ('Adaptability',
          'The system can run on different hosting environments (e.g., shared hosting, cloud/Docker) '
          'and databases (MySQL/PostgreSQL).*', True),
         ('Scalability',
          'The system can accommodate more barangays, users, and records as MENRO operations grow.*', True),
         ('Installability', 'The system can be deployed and configured by following the setup guide.*', True),
         ('Replaceability',
          'The system can replace the current manual/spreadsheet-based recording of MENRO.', False),
     ]),
    ('I. SAFETY',
     'Degree to which the system avoids states that endanger people, property, or the environment.',
     [
         ('Operational constraint',
          'Critical actions (deleting records, managing users) are restricted to authorized roles.', False),
         ('Risk identification',
          'The system highlights non-compliant generators, overdue follow-ups, and reported incidents.', False),
         ('Hazard warning',
          'Notifications promptly alert staff to issues that need action.', False),
         ('Fail safe',
          'When an error occurs, the system stops the action safely without saving incomplete data.', False),
     ]),
]

heading('PART III. EVALUATION')
W = [1.0, 3.6, 9.4, 0.7, 0.7, 0.7, 0.7, 0.7]
num = 0
for title, defn, items in SECTIONS:
    t = doc.add_table(rows=2, cols=8)
    borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Row 0: characteristic title across all columns
    top = t.rows[0].cells[0].merge(t.rows[0].cells[7])
    top.text = ''
    p = top.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p2 = top.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(defn)
    r2.italic = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade(top, GREEN)
    repeat_header(t.rows[0])

    # Row 1: column headers
    for i, h in enumerate(['No.', 'Sub-characteristic', 'Statement', '5', '4', '3', '2', '1']):
        c = t.rows[1].cells[i]
        cell_text(c, h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, LIGHT)
    repeat_header(t.rows[1])

    for sub, stmt, _tech in items:
        num += 1
        row = t.add_row()
        keep_row(row)
        cell_text(row.cells[0], str(num), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row.cells[1], sub, size=8.5, italic=True)
        cell_text(row.cells[2], stmt, size=9.5)
        for k in range(3, 8):
            cell_text(row.cells[k], '', size=9)
    set_widths(t, W)
    para(after=4)

# ── Part IV: Comments ────────────────────────────────────────────────────────
heading('PART IV. COMMENTS AND SUGGESTIONS')
for q in ('What features of the system did you find most useful?',
          'What problems did you encounter while using the system?',
          'What improvements would you recommend?'):
    para(q, after=2)
    for _ in range(2):
        p = para('_' * 95, size=9, after=2, color='9E9E9E')

para(after=10)
para('Thank you for taking the time to evaluate the MENRO Waste Management System.',
     italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

sig = doc.add_table(rows=2, cols=2)
cell_text(sig.rows[0].cells[1], '______________________________', align=WD_ALIGN_PARAGRAPH.CENTER)
cell_text(sig.rows[1].cells[1], 'Signature of Respondent', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
set_widths(sig, [9.0, 8.5])

out = 'MENRO_ISO25010_Evaluation_Questionnaire.docx'
doc.save(out)
print(f'Saved {out} with {num} items')
